import contextlib
import csv
import io
import json
import struct
import sys
import tempfile
import unittest
import zipfile
import zlib
from pathlib import Path

from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from ort import __version__
from ort.core.decoder import decode_local_file_data, recover_deflate_prefix, validate_xml_prefix
from ort.core.extractor import get_file_hex, recover_filename_length
from ort.core.fragment_salvage import recover_deflate_fragments
from ort.core.handlers.common import extract_xml_text_for_docx
from ort.core.utils import detect_image_format, validate_image_bytes, write_part
from ort.main import main


BODY = (
    b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    b'<w:body><w:p><w:r><w:t>Hel</w:t></w:r><w:r><w:t>lo</w:t></w:r>'
    b'<w:r><w:t xml:space="preserve"> world</w:t></w:r></w:p></w:body></w:document>'
)
FOOTER = (
    b'<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    b'<w:p><w:r><w:t>Footer text</w:t></w:r></w:p></w:ftr>'
)


def package(parts, compression=zipfile.ZIP_DEFLATED):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=compression) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        for name, data in parts.items():
            archive.writestr(name, data)
    return stream.getvalue()


def png_bytes():
    stream = io.BytesIO()
    Image.new("RGB", (2, 2), (120, 30, 200)).save(stream, format="PNG")
    return stream.getvalue()


class ApprovedRecoveryTests(unittest.TestCase):
    def test_run_boundaries_do_not_insert_spaces(self):
        self.assertEqual(extract_xml_text_for_docx(BODY.decode()), "Hello world")

    def test_stored_xml_uses_method_metadata(self):
        entries = [
            {
                "local_file_name": "word/document.xml",
                "local_file_data": BODY.hex(),
                "compression_method": 0,
            }
        ]
        decode_local_file_data(entries)
        self.assertEqual(entries[0]["decoded_xml"], BODY.decode())
        self.assertEqual(entries[0]["decode_status"], "stored")

    def test_prefix_validation_rejects_short_garbage(self):
        self.assertEqual(validate_xml_prefix(b"\r"), (False, "not_xml_start"))
        self.assertTrue(validate_xml_prefix(BODY[:180])[0])

    def test_linear_deflate_prefix(self):
        raw = zlib.compress(BODY)[2:-4]
        self.assertEqual(recover_deflate_prefix(raw), BODY)
        prefix = recover_deflate_prefix(raw[:-6])
        self.assertTrue(prefix and BODY.startswith(prefix))
        self.assertEqual(recover_deflate_prefix(b"\xff" * 1000), b"")

    def test_invalid_filename_length_is_bounded(self):
        part = bytearray(30) + b"word/document.xml" + b"\xff" * 4
        for declared in (0, 65535):
            struct.pack_into("<H", part, 26, declared)
            self.assertEqual(recover_filename_length(part), 17)

    def test_filename_length_does_not_absorb_deflate_bytes(self):
        name = b"word/webSettings.xml"
        part = bytearray(30) + name + b"rc" + b"payload"
        struct.pack_into("<H", part, 26, len(name) + 2)
        self.assertEqual(recover_filename_length(part), len(name))

    def test_root_relationship_part_name_is_valid(self):
        name = b"_rels/.rels"
        part = bytearray(30) + name + b"payload"
        struct.pack_into("<H", part, 26, len(name))
        self.assertEqual(recover_filename_length(part), len(name))

    def test_extensionless_ppt_metadata_part_name_is_valid(self):
        name = b"ppt/metadata"
        part = bytearray(30) + name + b"payload"
        struct.pack_into("<H", part, 26, len(name))
        self.assertEqual(recover_filename_length(part), len(name))

    def test_prefixed_package_is_processed_without_pseudo_entry(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "prefixed.docx"
            output_root = Path(temp_dir) / "results"
            input_path.write_bytes(b"J" * 64 + package({"word/document.xml": BODY}))
            with contextlib.redirect_stdout(io.StringIO()):
                result = get_file_hex(str(input_path), output_root=output_root, make_pdf=False)
            self.assertEqual(result.extension, ".docx")
            self.assertNotIn("J", "".join(item["local_file_name"] for item in result.entries))
            manifest = json.loads((output_root / "output_prefixed.docx" / "recovery_manifest.json").read_text())
            self.assertEqual(manifest["orphan_spans"][0]["kind"], "unowned_prefix")

    def test_footer_is_extracted_when_body_is_missing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "footer_only.docx"
            output_root = Path(temp_dir) / "results"
            input_path.write_bytes(package({"word/footer1.xml": FOOTER}))
            with contextlib.redirect_stdout(io.StringIO()):
                extension = get_file_hex(str(input_path), output_root=output_root, make_pdf=False).extension
            self.assertEqual(extension, ".docx")
            out = output_root / "output_footer_only.docx"
            manifest = json.loads((out / "recovery_manifest.json").read_text())
            self.assertEqual(manifest["inventory_status"], "complete")
            # v2: a missing core part is "missing" regardless of other recovered parts;
            # nothing was lost from this intact package, so overall is complete (Q2 strict rule).
            self.assertEqual(manifest["core_content_status"], "missing")
            self.assertEqual(manifest["document_content_status"], "complete")
            self.assertEqual(manifest["overall_status"], "complete")
            self.assertEqual(manifest["content_parts"][0]["name"], "word/footer1.xml")
            self.assertIn("Footer text", (out / "text" / "document_content.txt").read_text())
            self.assertFalse((out / "text" / "body_text.docx").exists())

    def test_complete_body_has_conservative_complete_status(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "clean.docx"
            output_root = Path(temp_dir) / "results"
            input_path.write_bytes(package({"word/document.xml": BODY}))
            with contextlib.redirect_stdout(io.StringIO()):
                get_file_hex(str(input_path), output_root=output_root, make_pdf=False)
            out = output_root / "output_clean.docx"
            manifest = json.loads((out / "recovery_manifest.json").read_text())
            self.assertEqual(manifest["inventory_status"], "complete")
            self.assertEqual(manifest["core_content_status"], "complete")
            self.assertEqual(manifest["image_status"], "none")
            self.assertEqual(manifest["overall_status"], "complete")
            self.assertTrue((out / "text" / "body_text.docx").exists())
            self.assertTrue((out / "text" / "word_document.xml.txt").exists())

    def test_png_crc_and_terminal_structure_are_checked(self):
        raw = png_bytes()
        self.assertEqual(detect_image_format(raw), "png")
        self.assertEqual(validate_image_bytes(raw, "png")[0], "decodable_structurally_complete")
        self.assertNotEqual(validate_image_bytes(raw[:-2], "png")[0], "decodable_structurally_complete")
        damaged = bytearray(raw)
        damaged[-16] ^= 1
        self.assertNotEqual(validate_image_bytes(bytes(damaged), "png")[0], "decodable_structurally_complete")

    def test_unverified_media_keeps_original_name_and_extension(self):
        emf = bytearray(64)
        emf[:4] = b"\x01\x00\x00\x00"
        emf[40:44] = b" EMF"
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "media.docx"
            output_root = Path(temp_dir) / "results"
            input_path.write_bytes(
                package({"word/document.xml": BODY, "word/media/original.emf": bytes(emf)})
            )
            with contextlib.redirect_stdout(io.StringIO()):
                get_file_hex(str(input_path), output_root=output_root, make_pdf=False)
            out = output_root / "output_media.docx"
            manifest = json.loads((out / "recovery_manifest.json").read_text())
            image = manifest["images"][0]
            # v2: integrity (CRC-32) and format validation are separate evidence fields.
            self.assertEqual(image["status"], "preserved_verified")
            self.assertEqual(image["integrity"], "crc_verified")
            self.assertEqual(image["format_validation"], "unsupported_format")
            self.assertTrue(image["output"].endswith("original.emf"))
            self.assertTrue((out / image["output"]).exists())
            self.assertEqual(manifest["image_status"], "all_preserved")
            self.assertEqual(manifest["overall_status"], "complete")

    def test_fragment_salvage_is_eof_and_evidence_gated(self):
        prefix = b'<w:document xmlns:w="x"><w:body><w:p><w:t>Lost</w:t></w:p>'
        suffix = b'<w:p><w:r><w:t>Recovered</w:t></w:r></w:p></w:body></w:document>'
        compressor = zlib.compressobj(level=6, wbits=-zlib.MAX_WBITS)
        stream = compressor.compress(prefix) + compressor.flush(zlib.Z_FULL_FLUSH)
        stream += compressor.compress(suffix) + compressor.flush(zlib.Z_FINISH)
        damaged = b"\x00" * 32 + stream[32:]
        recovered = list(recover_deflate_fragments(damaged, extension=".docx"))
        self.assertTrue(recovered)
        self.assertIn("Recovered", recovered[0]["text"])
        self.assertGreater(recovered[0]["non_whitespace_codepoints"], 0)
        self.assertEqual(list(recover_deflate_fragments(b"random" + b"\x00\x00\xff\xff" * 3, extension=".docx")), [])


# ----------------------------------------------------------------------------- v2 (F1–F6)


def local_entry(name, payload, method=0, crc=None, csize=None, usize=None, flags=0):
    raw_name = name.encode("utf-8")
    if not name.isascii():
        flags |= 0x800
    crc = zlib.crc32(payload) & 0xFFFFFFFF if crc is None else crc
    csize = len(payload) if csize is None else csize
    usize = len(payload) if usize is None else usize
    header = struct.pack("<4sHHHHHIIIHH", b"PK\x03\x04", 20, flags, method, 0, 0, crc, csize, usize,
                         len(raw_name), 0)
    return header + raw_name + payload


def entry_layout(data, name):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        info = archive.getinfo(name)
    name_length, extra_length = struct.unpack_from("<HH", data, info.header_offset + 26)
    return info.header_offset, info.header_offset + 30 + name_length + extra_length, info.compress_size


def run_tool(data, stem, salvage=False):
    temp_dir = tempfile.mkdtemp()
    input_path = Path(temp_dir) / f"{stem}.docx"
    input_path.write_bytes(data)
    with contextlib.redirect_stdout(io.StringIO()):
        get_file_hex(str(input_path), output_root=Path(temp_dir) / "results", salvage_fragments=salvage, make_pdf=False)
    out = Path(temp_dir) / "results" / f"output_{stem}.docx"
    return json.loads((out / "recovery_manifest.json").read_text(encoding="utf-8")), out


class StatusV2Tests(unittest.TestCase):
    def test_nonstandard_part_names_are_discovered(self):
        names = ["word/media/IMAGE_X.PNG", "word/media/photo.JPG", "word/media/image_x.emz",
                 "word/media/\uadf8\ub9bc1.png", "word/media/image(1).png", "word/media/media1.mp4",
                 "word/media/%EA%B7%B8%EB%A6%BC2.png"]
        parts = {"word/document.xml": BODY}
        parts.update({name: png_bytes() for name in names})
        manifest, _ = run_tool(package(parts), "names")
        found = {item["local_file_name"] for item in manifest["entries"]}
        self.assertTrue(set(names) <= found, set(names) - found)
        self.assertEqual(manifest["inventory_status"], "complete")
        self.assertEqual({image["status"] for image in manifest["images"]}, {"preserved_verified"})

    def test_filename_cannot_escape_or_contain_controls(self):
        for bad in (b"../evil.xml", b"/abs.xml", b"C:/x.xml", b"word/a\x01b.xml", b"word\\x.xml"):
            part = bytearray(30) + bad + b"\x00" * 8
            struct.pack_into("<H", part, 26, len(bad))
            self.assertNotEqual(recover_filename_length(bytes(part)), len(bad), bad)

    def test_stored_media_with_zeroed_sizes_uses_cd_then_next_header(self):
        image = png_bytes()
        original = package({"word/media/image1.png": image, "word/document.xml": BODY},
                           compression=zipfile.ZIP_STORED)
        offset, _, _ = entry_layout(original, "word/media/image1.png")
        damaged = bytearray(original)
        struct.pack_into("<II", damaged, offset + 18, 0, 0)
        without_cd = bytes(damaged[: bytes(damaged).find(b"PK\x01\x02")])
        for label, data, expected in (("with_cd", bytes(damaged), "central_directory"),
                                      ("without_cd", without_cd, "next_header_fallback")):
            manifest, out = run_tool(data, f"stored_{label}")
            record = next(item for item in manifest["images"] if item["name"] == "word/media/image1.png")
            self.assertEqual(record["status"], "preserved_verified", label)
            self.assertEqual((out / record["output"]).read_bytes(), image, label)
            entry = next(item for item in manifest["entries"] if item["local_file_name"] == "word/media/image1.png")
            self.assertEqual(entry["boundary_status"], expected, label)

    def test_stored_data_descriptor_boundary(self):
        image = png_bytes()
        descriptor = struct.pack("<4sIII", b"PK\x07\x08", zlib.crc32(image) & 0xFFFFFFFF, len(image), len(image))
        data = (local_entry("word/media/image1.png", image, crc=0, csize=0, usize=0, flags=0x8) + descriptor
                + local_entry("word/document.xml", BODY))
        manifest, _ = run_tool(data, "descriptor")
        entry = next(item for item in manifest["entries"] if item["local_file_name"] == "word/media/image1.png")
        self.assertEqual(entry["boundary_status"], "data_descriptor")
        self.assertEqual(entry["content_validation_status"], "crc_verified")
        self.assertEqual(manifest["inventory_status"], "unknown")

    def test_crc_mismatch_image_is_kept_as_incomplete(self):
        image = png_bytes()
        original = package({"word/document.xml": BODY, "word/media/image1.png": image},
                           compression=zipfile.ZIP_STORED)
        _, start, size = entry_layout(original, "word/media/image1.png")
        damaged = bytearray(original)
        damaged[start + size // 2] ^= 0xFF
        manifest, _ = run_tool(bytes(damaged), "crcmismatch")
        record = next(item for item in manifest["images"] if item["name"] == "word/media/image1.png")
        self.assertEqual(record["integrity"], "crc_mismatch")
        self.assertEqual(record["status"], "preserved_incomplete")
        self.assertEqual(manifest["overall_status"], "partial")

    def test_junk_complete_stream_is_rejected_and_not_written(self):
        image = png_bytes() * 50
        original = package({"word/document.xml": BODY, "word/media/image1.png": image})
        _, start, size = entry_layout(original, "word/media/image1.png")
        junk = zlib.compressobj(wbits=-zlib.MAX_WBITS)
        junk_stream = junk.compress(b"\x01") + junk.flush()
        damaged = original[:start] + junk_stream + original[start + size:]
        manifest, out = run_tool(damaged, "junk")
        record = next(item for item in manifest["images"] if item["name"] == "word/media/image1.png")
        self.assertEqual(record["status"], "rejected")
        self.assertIsNone(record["output"])
        self.assertEqual(record["rejected_bytes"], 1)
        self.assertIsNotNone(record["payload_offset"])
        self.assertFalse(list((out / "img").rglob("image1.png")))

    def test_rejected_stored_image_is_kept_as_diagnostic_bin(self):
        image = png_bytes() * 20
        original = package({"word/document.xml": BODY, "word/media/image1.png": image},
                           compression=zipfile.ZIP_STORED)
        _, start, size = entry_layout(original, "word/media/image1.png")
        damaged = bytearray(original)
        damaged[start:start + 64] = b"\x00" * 64  # signature destroyed, CRC mismatch
        manifest, out = run_tool(bytes(damaged), "storedhead")
        record = next(item for item in manifest["images"] if item["name"] == "word/media/image1.png")
        self.assertEqual(record["status"], "rejected")
        self.assertIsNone(record["output"])
        self.assertEqual(manifest["image_status"], "failed")
        self.assertTrue(record["diagnostic_output"].endswith("rejected_stored/image1.png.bin"))
        self.assertEqual((out / record["diagnostic_output"]).read_bytes(), bytes(damaged[start:start + size]))
        entry = next(item for item in manifest["entries"] if item["local_file_name"] == "word/media/image1.png")
        self.assertNotIn("rejected_stored_image_hex", entry)
        self.assertIsNone(record["uniform_fill_byte"])

    def test_uniformly_filled_stored_image_has_no_diagnostic_bin(self):
        image = png_bytes() * 20
        original = package({"word/document.xml": BODY, "word/media/image1.png": image},
                           compression=zipfile.ZIP_STORED)
        _, start, size = entry_layout(original, "word/media/image1.png")
        for fill in (0x00, 0xFF):
            damaged = bytearray(original)
            damaged[start:start + size] = bytes([fill]) * size  # the whole data area is one value
            manifest, out = run_tool(bytes(damaged), f"storedfill{fill:02x}")
            record = next(item for item in manifest["images"] if item["name"] == "word/media/image1.png")
            self.assertEqual(record["status"], "rejected")
            self.assertEqual(record["rejected_bytes"], size)
            self.assertEqual(record["uniform_fill_byte"], f"0x{fill:02X}")
            self.assertIsNone(record["diagnostic_output"])
            self.assertIsNone(record["diagnostic_sha256"])
            self.assertFalse((out / "img" / "rejected_stored").exists())

    def test_rejected_deflate_image_has_no_diagnostic_bin(self):
        image = png_bytes() * 50
        original = package({"word/document.xml": BODY, "word/media/image1.png": image})
        _, start, size = entry_layout(original, "word/media/image1.png")
        junk = zlib.compressobj(wbits=-zlib.MAX_WBITS)
        damaged = original[:start] + junk.compress(b"\x01") + junk.flush() + original[start + size:]
        manifest, out = run_tool(damaged, "deflatejunk")
        record = next(item for item in manifest["images"] if item["name"] == "word/media/image1.png")
        self.assertEqual(record["status"], "rejected")
        self.assertIsNone(record["diagnostic_output"])
        self.assertFalse((out / "img" / "rejected_stored").exists())

    # ------------------------------------------------------------------ F9 / v2.1

    def test_complete_stream_with_corrupt_middle_is_kept_as_crc_mismatch(self):
        tail = b"".join(b'<w:p><w:r><w:t>After %d</w:t></w:r></w:p>' % i for i in range(20))
        corrupt_body = BODY.replace(b"</w:body>", b"\xff\xfe\xfd" + tail + b"</w:body>")
        compressor = zlib.compressobj(wbits=-zlib.MAX_WBITS)
        stream = compressor.compress(corrupt_body) + compressor.flush()
        data = local_entry("word/document.xml", stream, method=8, crc=zlib.crc32(BODY), csize=len(stream),
                           usize=len(BODY))
        manifest, out = run_tool(data, "crcmiddle")
        entry = next(e for e in manifest["entries"] if e["local_file_name"] == "word/document.xml")
        self.assertEqual(entry["decode_status"], "complete")
        self.assertEqual(entry["content_validation_status"], "crc_mismatch")
        self.assertEqual(entry["utf8_replaced_sequences"], 3)
        self.assertEqual(entry["utf8_first_error_offset"], corrupt_body.index(b"\xff"))
        self.assertEqual(entry["decoded_bytes_sha256"], __import__("hashlib").sha256(corrupt_body).hexdigest())
        text = (out / "text" / "word_document.xml.txt").read_text(encoding="utf-8")
        self.assertIn("Hello world", text)
        self.assertIn("After 19", text)
        self.assertEqual(manifest["core_content_status"], "partial")

    def test_partial_prefix_with_split_trailing_character_is_trimmed(self):
        body = BODY.replace(b"Hel", "가나다".encode())
        cut = body.index("다".encode()) + 2  # stop inside the 3-byte character
        compressor = zlib.compressobj(0, zlib.DEFLATED, -zlib.MAX_WBITS)  # stored blocks: output is exact
        stream = compressor.compress(body) + compressor.flush()
        truncated = stream[: 5 + cut]  # 5-byte stored-block header, then the first `cut` bytes
        item = {"local_file_name": "word/document.xml", "compressed_data_hex": truncated.hex(), "compression_method": 8}
        decode_local_file_data([item])
        self.assertEqual(item["decode_status"], "partial_prefix")
        self.assertEqual(item["content_validation_status"], "prefix_validated")
        self.assertEqual(item["utf8_trailing_bytes_removed"], 2)
        self.assertEqual(item["decoded_xml"].encode(), body[: cut - 2])

    def test_xml_evidence_requires_namespace_and_complete_start_tag(self):
        namespace = b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        self.assertFalse(validate_xml_prefix(b"<w:document" + b"\xff" * 8, "word/document.xml")[0])
        self.assertFalse(validate_xml_prefix(b"<w:document " + namespace + b"\xff", "word/document.xml")[0])
        self.assertTrue(validate_xml_prefix(b"<w:document " + namespace + b">\xff\xfe", "word/document.xml")[0])
        self.assertFalse(validate_xml_prefix(b"<w:t>inserted marker</w:t>", "word/document.xml")[0])
        self.assertFalse(validate_xml_prefix(b"\x8a\x01<w:p>", "word/document.xml")[0])
        self.assertFalse(validate_xml_prefix(b"<a:p " + namespace + b">", "word/document.xml")[0])

    def test_garbage_complete_xml_without_evidence_stays_rejected(self):
        junk = b"<w:document" + bytes(range(128, 256))
        compressor = zlib.compressobj(wbits=-zlib.MAX_WBITS)
        stream = compressor.compress(junk) + compressor.flush()
        data = local_entry("word/document.xml", stream, method=8, crc=zlib.crc32(BODY), csize=len(stream))
        manifest, _ = run_tool(data, "junkxml")
        entry = next(e for e in manifest["entries"] if e["local_file_name"] == "word/document.xml")
        self.assertEqual(entry["content_validation_status"], "rejected")
        self.assertEqual(manifest["core_content_status"], "failed")

    def test_unverified_stream_end_does_not_skip_embedded_header(self):
        # A DEFLATE stream whose end lies beyond a real local header (stored blocks keep
        # that header literal). Its CRC does not verify and its size is unknown, so the
        # scanner must not jump over the embedded footer entry.
        footer_entry = local_entry("word/footer1.xml", FOOTER)
        compressor = zlib.compressobj(0, zlib.DEFLATED, -zlib.MAX_WBITS)
        stream = compressor.compress(b"<x/>" + footer_entry + b"<y/>") + compressor.flush()
        data = (local_entry("word/document.xml", stream, method=8, crc=zlib.crc32(BODY), csize=0, usize=len(BODY))
                + local_entry("word/endnotes.xml", FOOTER))
        manifest, out = run_tool(data, "falseeof")
        names = [item["local_file_name"] for item in manifest["entries"]]
        self.assertIn("word/footer1.xml", names)
        self.assertIn("Footer text", (out / "text" / "word_footer1.xml.txt").read_text(encoding="utf-8"))

    def test_verified_stored_package_hides_inner_headers(self):
        inner = package({"xl/worksheets/sheet1.xml": b"<worksheet/>"})
        data = package({"word/document.xml": BODY, "word/embeddings/Book1.xlsx": inner},
                       compression=zipfile.ZIP_STORED)
        manifest, _ = run_tool(data, "embedded")
        names = [item["local_file_name"] for item in manifest["entries"]]
        self.assertNotIn("xl/worksheets/sheet1.xml", names)
        self.assertEqual(manifest["inventory_status"], "complete")

    def test_overall_is_strict_and_reports_auxiliary_failures(self):
        styles = b"<w:styles xmlns:w='x'>" + b"<a/>" * 200 + b"</w:styles>"
        original = package({"word/document.xml": BODY, "word/styles.xml": styles})
        _, start, _ = entry_layout(original, "word/styles.xml")
        damaged = bytearray(original)
        damaged[start:start + 8] = b"\x00" * 8
        manifest, _ = run_tool(bytes(damaged), "strict")
        self.assertEqual(manifest["core_content_status"], "complete")
        self.assertEqual(manifest["document_content_status"], "complete")
        self.assertEqual(manifest["overall_status"], "partial")
        self.assertEqual(manifest["auxiliary_failed_entries"], 1)
        self.assertEqual(manifest["content_failed_entries"], 0)

    def test_xlsx_and_pptx_document_content_status(self):
        sheet = (b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
                 b'<row r="1"><c r="A1" t="s"><v>0</v></c></row></sheetData></worksheet>')
        shared = b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>Cell</t></si></sst>'
        manifest, _ = run_tool(package({"xl/worksheets/sheet1.xml": sheet, "xl/sharedStrings.xml": shared}), "book")
        self.assertEqual(manifest["inferred_type"], ".xlsx")
        self.assertEqual(manifest["document_content_status"], "complete")
        slide = (b'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                 b'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree><p:sp><p:txBody>'
                 b'<a:p><a:r><a:t>Slide</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
        notes = slide.replace(b"p:sld", b"p:notes").replace(b"Slide", b"Speaker note")
        manifest, out = run_tool(package({"ppt/slides/slide1.xml": slide, "ppt/notesSlides/notesSlide1.xml": notes}),
                                 "deck")
        self.assertEqual(manifest["document_content_status"], "complete")
        self.assertIn("Speaker note", (out / "text" / "notesSlide1.txt").read_text(encoding="utf-8"))

    def test_manifest_records_schema_and_code_hash(self):
        manifest, _ = run_tool(package({"word/document.xml": BODY}), "hash")
        self.assertEqual(manifest["schema_version"], 3)
        self.assertRegex(manifest["tool_code_sha256"], r"^[0-9a-f]{64}$")


# ----------------------------------------------------------------------------- 0.3.0 (A1–C5)

SSML = b'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
DML = b'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
PML = b'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'


def read_csv_rows(path):
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    return list(csv.reader(line for line in lines if not line.startswith("#")))


def run_named(parts, file_name, **options):
    temp_dir = Path(tempfile.mkdtemp())
    input_path = temp_dir / file_name
    input_path.write_bytes(package(parts))
    with contextlib.redirect_stdout(io.StringIO()):
        options.setdefault("make_pdf", False)  # PDF conversion is tested separately
        result = get_file_hex(str(input_path), output_root=temp_dir / "results", **options)
    manifest = json.loads((result.output_directory / "recovery_manifest.json").read_text(encoding="utf-8"))
    return manifest, result.output_directory


def truncated_body_entry():
    body = BODY.replace(b"</w:body>", b"".join(b"<w:p><w:r><w:t>Line %d</w:t></w:r></w:p>" % i
                                               for i in range(50)) + b"</w:body>")
    compressor = zlib.compressobj(0, zlib.DEFLATED, -zlib.MAX_WBITS)  # stored blocks: the prefix is exact
    stream = compressor.compress(body) + compressor.flush()
    truncated = stream[: len(stream) // 2]
    return local_entry("word/document.xml", truncated, method=8, crc=zlib.crc32(body), csize=len(truncated),
                       usize=len(body))


class Release030Tests(unittest.TestCase):
    def test_shared_strings_follow_si_items(self):  # A1
        shared = (b'<sst ' + SSML + b'><si><t>alpha</t></si><si><t xml:space="preserve"> beta</t></si>'
                  b'<si><r><t>gam</t></r><r><rPr><b/></rPr><t>ma</t></r><rPh sb="0" eb="1"><t>GA</t></rPh></si>'
                  b'<si><t>delta</t></si></sst>')
        sheet = (b'<worksheet ' + SSML + b'><sheetData><row r="1">'
                 b'<c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>2</v></c>'
                 b'<c r="D1" t="s"><v>3</v></c><c r="E1" t="s"><v>9</v></c>'
                 b'<c r="F1" t="inlineStr"><is><r><t>in</t></r><r><t>line</t></r></is></c>'
                 b'<c r="AA1" s="1"><v>45000</v></c></row>'
                 b'<row r="2"><c r="B2"><f>SUM(1,2)</f><v>3</v></c></row></sheetData></worksheet>')
        styles = (b'<styleSheet ' + SSML + b'><cellStyleXfs count="1"><xf numFmtId="4"/></cellStyleXfs>'
                  b'<cellXfs count="2"><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs></styleSheet>')
        _, out = run_named({"xl/worksheets/sheet1.xml": sheet, "xl/sharedStrings.xml": shared,
                            "xl/styles.xml": styles}, "book.xlsx")
        grid = read_csv_rows(out / "text" / "sheet1.csv")
        self.assertEqual(grid[0], ["Row", "A", "B", "C", "D", "E", "F", "AA"])
        self.assertEqual(grid[1], ["1", "alpha", " beta", "gamma", "delta", "[shared string #9 not recovered]",
                                   "inline", "45000"])
        cells = {row[0]: row for row in read_csv_rows(out / "text" / "sheet1_cells.csv")[1:]}
        self.assertEqual(cells["AA1"][3:], ["45000", "", "mm-dd-yy"])
        self.assertEqual(cells["B2"][3:], ["3", "SUM(1,2)", "General"])

    def test_sheet_files_use_part_number_and_display_name(self):  # A2, B4
        sheet = b'<worksheet ' + SSML + b'><sheetData><row r="1"><c r="A1"><v>333</v></c></row></sheetData></worksheet>'
        workbook = (b'<workbook ' + SSML + b' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                    b'<sheets><sheet name="Sales" sheetId="3" r:id="rId3"/></sheets></workbook>')
        rels = (b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                b'<Relationship Id="rId3" Target="worksheets/sheet3.xml" '
                b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/></Relationships>')
        manifest, out = run_named({"xl/worksheets/sheet3.xml": sheet, "xl/workbook.xml": workbook,
                                   "xl/_rels/workbook.xml.rels": rels}, "book.xlsx")
        self.assertFalse((out / "text" / "sheet1.csv").exists())
        self.assertEqual(read_csv_rows(out / "text" / "sheet3.csv")[1], ["1", "333"])
        header = (out / "text" / "sheet3.csv").read_text(encoding="utf-8-sig").splitlines()[:2]
        self.assertTrue(header[0].startswith("# source: xl/worksheets/sheet3.xml | crc_verified"))
        self.assertEqual(header[1], "# sheet name: Sales (from xl/workbook.xml)")
        record = next(part for part in manifest["content_parts"] if part["name"] == "xl/worksheets/sheet3.xml")
        self.assertEqual(record["sheet_name"], "Sales")
        self.assertIn("rId3: worksheets/sheet3.xml", (out / "relationships.txt").read_text(encoding="utf-8"))

    def test_output_folder_keeps_extension_and_is_never_reused(self):  # A3
        temp_dir = Path(tempfile.mkdtemp())
        for extension, parts in ((".docx", {"word/document.xml": BODY}),
                                 (".xlsx", {"xl/sharedStrings.xml": b"<sst/>"}),
                                 (".pptx", {"ppt/slides/slide1.xml": b"<p:sld/>"})):
            (temp_dir / f"same{extension}").write_bytes(package(parts))
        with contextlib.redirect_stdout(io.StringIO()):
            folders = [get_file_hex(str(temp_dir / name), output_root=temp_dir / "out", make_pdf=False).output_directory.name
                       for name in ("same.docx", "same.xlsx", "same.pptx", "same.docx")]
        self.assertEqual(folders, ["output_same.docx", "output_same.xlsx", "output_same.pptx", "output_same.docx_2"])

    def test_validated_prefix_is_flagged_in_manifest_and_text(self):  # A4
        manifest, out = run_tool(truncated_body_entry(), "prefix")
        entry = next(e for e in manifest["entries"] if e["local_file_name"] == "word/document.xml")
        self.assertEqual(entry["content_validation_status"], "prefix_validated")
        self.assertTrue(entry["tail_may_be_non_original"])
        first_line = (out / "text" / "word_document.xml.txt").read_text(encoding="utf-8").splitlines()[0]
        self.assertIn("prefix_validated", first_line)
        self.assertIn("may contain non-original text", first_line)

    def test_pptx_text_keeps_runs_paragraphs_tables_and_groups(self):  # A5
        slide = (b'<p:sld ' + PML + b' ' + DML + b'><p:cSld><p:spTree>'
                 b'<p:sp><p:nvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr><p:txBody>'
                 b'<a:p><a:r><a:t>Forens</a:t></a:r><a:r><a:t>ic</a:t></a:r></a:p></p:txBody></p:sp>'
                 b'<p:sp><p:txBody><a:p><a:r><a:t>First</a:t></a:r><a:br/><a:r><a:t>break</a:t></a:r></a:p>'
                 b'<a:p><a:r><a:t>Second</a:t></a:r></a:p></p:txBody></p:sp>'
                 b'<p:grpSp><p:sp><p:txBody><a:p><a:r><a:t>Grouped</a:t></a:r></a:p></p:txBody></p:sp></p:grpSp>'
                 b'<p:graphicFrame><a:graphic><a:graphicData><a:tbl><a:tr>'
                 b'<a:tc><a:txBody><a:p><a:r><a:t>R1C1</a:t></a:r></a:p></a:txBody></a:tc>'
                 b'<a:tc><a:txBody><a:p><a:r><a:t>R1C2</a:t></a:r></a:p></a:txBody></a:tc>'
                 b'</a:tr></a:tbl></a:graphicData></a:graphic></p:graphicFrame>'
                 b'</p:spTree></p:cSld></p:sld>')
        _, out = run_named({"ppt/slides/slide2.xml": slide}, "deck.pptx")
        lines = (out / "text" / "slide2.txt").read_text(encoding="utf-8").splitlines()[1:]
        self.assertEqual(lines, ["[TITLE]", "Forensic", "[BODY]", "First", "break", "Second", "[BODY]", "Grouped",
                                 "[TABLE]", "R1C1\tR1C2"])
        self.assertFalse((out / "text" / "slide1.txt").exists())

    def test_parts_are_written_byte_identical_and_entries_json_is_optional(self):  # B1, B2
        styles = b'<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
        parts = {"word/document.xml": BODY, "word/styles.xml": styles,
                 "word/embeddings/Book1.xlsx": package({"xl/worksheets/sheet1.xml": b"<worksheet/>"}),
                 "word/media/image1.png": png_bytes()}
        manifest, out = run_named(parts, "parts.docx")
        for name, data in parts.items():
            entry = next(e for e in manifest["entries"] if e["local_file_name"] == name)
            if "/media/" in name:
                self.assertNotIn("output", entry)
                continue
            self.assertEqual(entry["output"], "parts/" + name)
            self.assertEqual((out / entry["output"]).read_bytes(), data, name)
        self.assertEqual(list(out.glob("*.json")), [out / "recovery_manifest.json"])
        _, out = run_named(parts, "parts.docx", dump_entries=True)
        dumped = json.loads((out / "entries.json").read_text(encoding="utf-8"))
        self.assertIn("compressed_data_hex", dumped[0])

    def test_part_paths_are_safe_on_windows(self):  # B2
        root, used = Path(tempfile.mkdtemp()), set()
        written = [write_part(root, name, data, used) for name, data in
                   (("word/CON.xml", b"1"), ("word/Document.xml", b"2"), ("word/document.xml", b"3"))]
        self.assertEqual([path.name for path in written], ["_CON.xml", "Document.xml", "document_2.xml"])
        self.assertEqual([path.read_bytes() for path in written], [b"1", b"2", b"3"])

    def test_metadata_is_written_once_from_core_and_app(self):  # B3
        core = (b'<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
                b'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:creator>Alice</dc:creator>'
                b'<cp:revision>4</cp:revision></cp:coreProperties>')
        app = (b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
               b'<Application>Microsoft Office PowerPoint</Application><TotalTime>12</TotalTime></Properties>')
        slide = b'<p:sld ' + PML + b' ' + DML + b'><p:cSld><p:spTree/></p:cSld></p:sld>'
        manifest, out = run_named({"ppt/slides/slide1.xml": slide, "docProps/core.xml": core,
                                   "docProps/app.xml": app}, "meta.pptx")
        text = (out / "metadata.txt").read_text(encoding="utf-8")
        self.assertIn("# source: docProps/core.xml | crc_verified", text)
        self.assertIn("creator: Alice", text)
        self.assertIn("Application: Microsoft Office PowerPoint", text)
        self.assertFalse((out / "metadata").exists())
        self.assertEqual(manifest["metadata"]["docProps/app.xml"]["values"]["TotalTime"], "12")

    def test_manifest_v3_fields(self):  # B5, C1
        manifest, _ = run_tool(package({"word/document.xml": BODY}), "v3")
        self.assertEqual(manifest["output_layout"], 2)
        self.assertEqual(manifest["tool_version"], __version__)
        self.assertEqual(manifest["evaluation_definitions"], "evaluation_definitions_v2_1")
        self.assertEqual(manifest["entries"][0]["extension"], ".xml")
        self.assertNotIn("img_ext", manifest["entries"][0])

    def test_cli_exit_status_version_and_recursion(self):  # C2
        temp_dir = Path(tempfile.mkdtemp())
        (temp_dir / "sub").mkdir()
        (temp_dir / "sub" / "good.docx").write_bytes(package({"word/document.xml": BODY}))
        (temp_dir / "partial.docx").write_bytes(truncated_body_entry())
        (temp_dir / "notes.txt").write_text("not a package")
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(main([str(temp_dir / "sub" / "good.docx"), "--no-pdf"]), 0)
            self.assertEqual(main([str(temp_dir / "partial.docx"), "--no-pdf"]), 1)
            self.assertEqual(main([str(temp_dir), "--no-pdf"]), 2)  # notes.txt is not OOXML
            self.assertEqual(main([str(temp_dir / "sub"), "-r", "--no-pdf"]), 0)  # output_good.docx is skipped
            self.assertEqual(main([str(temp_dir / "missing"), "--no-pdf"]), 2)
            with self.assertRaises(SystemExit) as raised:
                main(["--version"])
        self.assertEqual(raised.exception.code, 0)
        self.assertFalse((temp_dir / "sub" / "output_good.docx" / "output_output_good.docx").exists())

    def test_max_part_size_limits_decoding(self):  # C5
        manifest, _ = run_named({"word/document.xml": BODY}, "small.docx", max_part_size=64)
        entry = next(e for e in manifest["entries"] if e["local_file_name"] == "word/document.xml")
        self.assertEqual(entry["decode_error"], "decoded output limit exceeded")
        self.assertEqual(manifest["max_part_size_bytes"], 64)

# ----------------------------------------------------------------------------- reading outputs

RNS = b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


def rels(*items):
    body = "".join(f'<Relationship Id="{rid}" Target="{target}" Type="{REL}{kind}"/>' for rid, target, kind in items)
    return f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{body}</Relationships>'.encode()


class ReadingOutputTests(unittest.TestCase):
    """recovered_report.html, the rebuilt document, and the PDF preview field."""

    DOC_WITH_IMAGE = (b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" ' + DML
                      + b' ' + RNS + b'><w:body><w:p><w:r><w:t>Before</w:t></w:r><w:r><w:drawing><a:graphic>'
                      b'<a:graphicData><a:blip r:embed="rId5"/></a:graphicData></a:graphic></w:drawing></w:r>'
                      b'<w:r><w:t>After</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>R1C1</w:t></w:r></w:p></w:tc>'
                      b'<w:tc><w:p><w:r><w:t>R1C2</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:p><w:r><w:t>R2C1</w:t>'
                      b'</w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>R2C2</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>')

    def docx_parts(self):
        return {"word/document.xml": self.DOC_WITH_IMAGE, "word/media/image1.png": png_bytes(),
                "word/_rels/document.xml.rels": rels(("rId5", "media/image1.png", "image"))}

    def xlsx_parts(self):
        sheet = (b'<worksheet ' + SSML + b' ' + RNS + b'><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c>'
                 b'<c r="B1"><v>42</v></c></row></sheetData><drawing r:id="rId1"/></worksheet>')
        drawing = (b'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                   + DML + b' ' + RNS + b'><xdr:twoCellAnchor><xdr:from><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff>'
                   b'<xdr:row>4</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:pic><xdr:blipFill>'
                   b'<a:blip r:embed="rId1"/></xdr:blipFill></xdr:pic></xdr:twoCellAnchor></xdr:wsDr>')
        workbook = (b'<workbook ' + SSML + b' ' + RNS + b'><sheets><sheet name="Sales" sheetId="1" r:id="rId1"/>'
                    b'</sheets></workbook>')
        return {"xl/workbook.xml": workbook, "xl/_rels/workbook.xml.rels": rels(("rId1", "worksheets/sheet1.xml", "worksheet")),
                "xl/worksheets/sheet1.xml": sheet, "xl/sharedStrings.xml": b'<sst ' + SSML + b'><si><t>Cell</t></si></sst>',
                "xl/worksheets/_rels/sheet1.xml.rels": rels(("rId1", "../drawings/drawing1.xml", "drawing")),
                "xl/drawings/drawing1.xml": drawing,
                "xl/drawings/_rels/drawing1.xml.rels": rels(("rId1", "../media/image1.png", "image")),
                "xl/media/image1.png": png_bytes()}

    def pptx_parts(self):
        presentation = (b'<p:presentation ' + PML + b' ' + RNS + b'><p:sldIdLst><p:sldId id="256" r:id="rId2"/>'
                        b'</p:sldIdLst></p:presentation>')
        slide = (b'<p:sld ' + PML + b' ' + DML + b' ' + RNS + b'><p:cSld><p:spTree><p:sp><p:spPr><a:xfrm>'
                 b'<a:off x="457200" y="457200"/><a:ext cx="4572000" cy="914400"/></a:xfrm></p:spPr><p:txBody><a:p><a:r>'
                 b'<a:t>Title text</a:t></a:r></a:p></p:txBody></p:sp><p:pic><p:blipFill><a:blip r:embed="rId3"/>'
                 b'</p:blipFill><p:spPr><a:xfrm><a:off x="457200" y="1828800"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
                 b'</p:spPr></p:pic></p:spTree></p:cSld></p:sld>')
        notes = (b'<p:notes ' + PML + b' ' + DML + b'><p:cSld><p:spTree><p:sp><p:nvSpPr><p:nvPr><p:ph type="body"/>'
                 b'</p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>Speaker note</a:t></a:r></a:p></p:txBody></p:sp>'
                 b'</p:spTree></p:cSld></p:notes>')
        return {"ppt/presentation.xml": presentation,
                "ppt/_rels/presentation.xml.rels": rels(("rId2", "slides/slide2.xml", "slide")),
                "ppt/slides/slide2.xml": slide, "ppt/notesSlides/notesSlide1.xml": notes,
                "ppt/slides/_rels/slide2.xml.rels": rels(("rId3", "../media/image1.png", "image"),
                                                         ("rId4", "../notesSlides/notesSlide1.xml", "notesSlide")),
                "ppt/media/image1.png": png_bytes()}

    def html(self, out):
        return (out / "recovered_report.html").read_text(encoding="utf-8")

    # ------------------------------------------------------------------ recovered_report.html
    def test_html_docx_image_position_table_and_label_index(self):
        manifest, out = run_named(self.docx_parts(), "doc.docx")
        self.assertEqual(manifest["recovered_report"], {"path": "recovered_report.html", "error": None})
        self.assertFalse((out / "recovered_document.md").exists())
        page = self.html(out)
        self.assertIn('Before<span class="mark">Image 1</span>After', page)
        self.assertIn('<img src="img/verified/image1.png"', page)
        self.assertEqual(page.count('<table class="doc">'), 1)  # consecutive rows form one table
        self.assertIn("R2C2", page)
        labels = page.split('id="labels"', 1)[1]
        self.assertIn("word/document.xml: “Body”", labels)
        self.assertNotIn("<script", page.lower())

    def test_html_guide_explains_every_label_below_the_header(self):
        from ort.core.document import LABEL_GUIDE, LABEL_ORDER
        _, out = run_named({"word/document.xml": BODY}, "guide.docx")
        page = self.html(out)
        self.assertLess(page.index('class="hero"'), page.index('id="guide"'))
        self.assertLess(page.index('id="guide"'), page.index('id="summary"'))
        guide = page.split('id="guide"', 1)[1].split('id="summary"', 1)[0]
        for status in LABEL_ORDER:
            self.assertIn(f">{status}</span>", guide)
            self.assertIn(LABEL_GUIDE[status][0], guide)
        for overall in ("complete", "partial", "failed", "indeterminate"):
            self.assertIn(f"<b>{overall}</b>", guide)
        self.assertIn('<a href="#guide">Guide</a>', page)

    def test_html_escapes_recovered_text(self):
        document = BODY.replace(b"Hel", b"&lt;script&gt;alert(1)&lt;/script&gt;Hel")
        _, out = run_named({"word/document.xml": document}, "escape.docx")
        page = self.html(out)
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", page)
        self.assertNotIn("<script>alert", page)

    def test_html_marks_the_damage_point_after_a_prefix(self):
        _, out = run_tool(truncated_body_entry(), "prefixdoc")
        page = self.html(out)
        article = page.split('id="part-1"', 1)[1].split("</article>", 1)[0]
        self.assertIn("prefix_validated", article)
        self.assertTrue(article.rstrip().endswith("the text just above may be non-original.</div>\n</div>"))

    def test_html_xlsx_grid_and_anchored_image(self):
        _, out = run_named(self.xlsx_parts(), "book.xlsx")
        page = self.html(out)
        self.assertIn("Sheet 1 &#x27;Sales&#x27;", page)
        self.assertIn("<th>A</th><th>B</th>", page)
        self.assertIn('<td class="rn">1</td><td>Cell</td><td>42</td>', page)
        self.assertIn("At C5: ", page)

    def test_html_pptx_order_picture_and_notes(self):
        _, out = run_named(self.pptx_parts(), "deck.pptx")
        page = self.html(out)
        self.assertIn("<h3>Slide 1</h3><span class=\"path\">ppt/slides/slide2.xml</span>", page)
        self.assertLess(page.index("Title text"), page.index('<span class="mark">Image 1</span>'))
        self.assertIn("<h3>Slide 1 notes</h3>", page)

    def test_html_describes_a_uniformly_filled_image(self):
        original = package({"word/document.xml": BODY, "word/media/image1.png": png_bytes() * 20},
                           compression=zipfile.ZIP_STORED)
        _, start, size = entry_layout(original, "word/media/image1.png")
        _, out = run_tool(original[:start] + b"\x00" * size + original[start + size:], "zerodoc")
        self.assertIn("its data area is filled with 0x00", self.html(out))

    def test_deeply_nested_xml_does_not_exceed_the_recursion_limit(self):
        depth = 3000  # deeper than Python's default recursion limit (1000)
        nested = b"<w:r>" * depth + b"<w:t>deep text</w:t>" + b"</w:r>" * depth
        document = (b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                    b"<w:sdt><w:sdtContent><w:p>" + nested + b"</w:p></w:sdtContent></w:sdt></w:body></w:document>")
        manifest, out = run_named({"word/document.xml": document}, "deep.docx")
        self.assertIsNone(manifest["recovered_report"]["error"])
        self.assertIsNone(manifest["rebuilt_document"]["error"])
        self.assertIn("deep text", self.html(out))

    def test_a_failing_reading_output_does_not_affect_recovery(self):
        from unittest import mock
        with mock.patch("ort.core.document.render_html", side_effect=RuntimeError("boom")):
            manifest, out = run_tool(package({"word/document.xml": BODY}), "isolated")
        self.assertEqual(manifest["overall_status"], "complete")
        self.assertEqual(manifest["recovered_report"], {"path": None, "error": "RuntimeError: boom"})
        self.assertFalse((out / "recovered_report.html").exists())
        self.assertTrue((out / manifest["rebuilt_document"]["path"]).exists())
        self.assertTrue((out / "text" / "word_document.xml.txt").exists())

    # ------------------------------------------------------------------ rebuilt document
    def test_rebuilt_docx_has_banner_content_table_image_and_labels(self):
        from docx import Document
        manifest, out = run_named(self.docx_parts(), "doc.docx")
        self.assertEqual(manifest["rebuilt_document"], {"path": "recovered_doc.docx", "error": None})
        document = Document(str(out / "recovered_doc.docx"))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertTrue(text.startswith("Rebuilt by ORT from the recovered content. This is NOT the original document"))
        self.assertIn("BeforeAfter", text.replace("[Image 1]", ""))
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(1, 1).text, "R2C2")
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertIn("ORT labels (index by validation status)", text)

    def test_rebuilt_xlsx_keeps_cells_names_formats_and_marks_damage(self):
        import openpyxl
        _, out = run_named(self.xlsx_parts(), "book.xlsx")
        workbook = openpyxl.load_workbook(str(out / "recovered_book.xlsx"))
        self.assertEqual(workbook.sheetnames, ["ORT recovery", "Sales"])
        self.assertTrue(str(workbook["ORT recovery"]["A1"].value).startswith("Rebuilt by ORT"))
        sheet = workbook["Sales"]
        self.assertEqual((sheet["A1"].value, sheet["B1"].value), ("Cell", 42))
        self.assertEqual(len(sheet._images), 1)
        self.assertIsNone(sheet.sheet_properties.tabColor)  # verified sheet: no warning colour

    def test_rebuilt_pptx_keeps_order_positions_pictures_and_notes(self):
        from pptx import Presentation
        _, out = run_named(self.pptx_parts(), "deck.pptx")
        deck = Presentation(str(out / "recovered_deck.pptx"))
        slides = list(deck.slides)
        self.assertIn("not the original", slides[0].shapes[0].text_frame.text)
        recovered = slides[1]
        texts = [shape.text_frame.text for shape in recovered.shapes if shape.has_text_frame]
        self.assertIn("Title text", texts)
        title = next(shape for shape in recovered.shapes if shape.has_text_frame and shape.text_frame.text == "Title text")
        self.assertEqual((title.left, title.top), (457200, 457200))
        self.assertEqual(sum(1 for shape in recovered.shapes if shape.shape_type == 13), 1)  # one picture
        self.assertIn("Speaker note", recovered.notes_slide.notes_text_frame.text)

    # ------------------------------------------------------------------ PDF preview
    def test_pdf_disabled_records_how_to_obtain_it(self):
        manifest, _ = run_named({"word/document.xml": BODY}, "nopdf.docx")
        pdf = manifest["preview_pdf"]
        self.assertEqual((pdf["status"], pdf["path"], pdf["reason"]), ("disabled", None, "disabled by --no-pdf"))
        self.assertEqual(pdf["how_to"][0], "Run ORT again without --no-pdf.")
        self.assertIn("open recovered_nopdf.docx in Word, Excel, PowerPoint, or LibreOffice and save it as PDF.",
                      pdf["how_to"])

    def test_pdf_skipped_without_a_converter(self):
        from unittest import mock
        with mock.patch("ort.core.pdf.find_soffice", return_value=None), mock.patch("ort.core.pdf.sys.platform", "linux"):
            manifest, out = run_named({"word/document.xml": BODY}, "skip.docx", make_pdf=True)
        pdf = manifest["preview_pdf"]
        self.assertEqual(pdf["status"], "skipped")
        self.assertIn("no PDF converter found", pdf["reason"])
        self.assertEqual(pdf["how_to"][0], "Install LibreOffice (with 'soffice' on PATH) and run ORT again, or")
        self.assertIn("PDF preview: skipped", self.html(out))

    def test_pdf_written_by_a_converter_is_recorded_and_linked(self):
        from unittest import mock

        def fake(soffice, source):
            source.with_suffix(".pdf").write_bytes(b"%PDF-1.4\n")
            return "written", None
        with mock.patch("ort.core.pdf.find_soffice", return_value="soffice"), mock.patch("ort.core.pdf._with_soffice", fake):
            manifest, out = run_named({"word/document.xml": BODY}, "made.docx", make_pdf=True)
        self.assertEqual(manifest["preview_pdf"]["status"], "written")
        self.assertEqual(manifest["preview_pdf"]["converter"], "LibreOffice")
        self.assertEqual(manifest["preview_pdf"]["path"], "recovered_made.pdf")
        self.assertIsNone(manifest["preview_pdf"]["how_to"])
        self.assertIn('href="recovered_made.pdf"', self.html(out))

    def test_cli_no_pdf_option(self):
        temp_dir = Path(tempfile.mkdtemp())
        (temp_dir / "cli.docx").write_bytes(package({"word/document.xml": BODY}))
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(main([str(temp_dir / "cli.docx"), "--no-pdf", "--no-pdf"]), 0)
        manifest = json.loads((temp_dir / "output_cli.docx" / "recovery_manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["preview_pdf"]["status"], "disabled")


if __name__ == "__main__":
    unittest.main()
